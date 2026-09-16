"""Read-only team-roster column inventory for the known annual workbook.

The existing Worship workbook parser remains authoritative for workbook shape,
supported rows, dates, and rotation tokens.  Its exact ServiceProfile and
ServiceEvent matcher remains authoritative for CMS target evidence.  This
adapter only adds literal header discovery and optional exact-header hints.

Columns A and B are returned as inventory evidence but are explicitly structural:
A owns event-date identity and B owns the Worship rotation token.  Columns C
through I are review candidates for the exact Bethany target profile.  Columns
J through O remain visible parallel-service evidence but are not mappable by
this adapter.  A blank XLSX header cell is represented by the empty string; no
nonblank header is trimmed, normalized, or otherwise rewritten.
"""

from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ministry.services.team_roster_workbook import (
    ObservedTeamRosterColumn,
    TeamRosterColumnHint,
    resolve_exact_team_roster_column_hint,
    validate_team_roster_column_hints,
)
from ministry.services.worship_xlsx_preview import (
    CONTRACT_REVISION,
    INTEGRATION_KEY,
    SUPPORTED_SHEET,
    ParsedWorshipWorkbook,
    TargetMatch,
    match_exact_service_event_targets,
    parse_known_worship_workbook,
)


TEAM_ROSTER_HEADER_ROW = 3
TEAM_ROSTER_INVENTORY_LAST_COLUMN = 15  # O, the parser's supported A1:O58 region.


class TeamRosterColumnScope(StrEnum):
    STRUCTURAL_EVENT_IDENTITY = "structural_event_identity"
    STRUCTURAL_WORSHIP_IDENTITY = "structural_worship_identity"
    TARGET_PROFILE_REVIEW_CANDIDATE = "target_profile_review_candidate"
    OUT_OF_TARGET_PROFILE_SCOPE = "out_of_target_profile_scope"


TEAM_ROSTER_COLUMN_SCOPE = {
    "A": TeamRosterColumnScope.STRUCTURAL_EVENT_IDENTITY,
    "B": TeamRosterColumnScope.STRUCTURAL_WORSHIP_IDENTITY,
    **{
        column: TeamRosterColumnScope.TARGET_PROFILE_REVIEW_CANDIDATE
        for column in "CDEFGHI"
    },
    **{
        column: TeamRosterColumnScope.OUT_OF_TARGET_PROFILE_SCOPE
        for column in "JKLMNO"
    },
}

TEAM_ROSTER_COLUMN_HINTS = validate_team_roster_column_hints(
    (
        TeamRosterColumnHint(
            expected_header="projector",
            destination_team_key="main.cm.digital.projection",
        ),
        TeamRosterColumnHint(
            expected_header="Sound",
            destination_team_key="main.cm.digital.sound",
        ),
        TeamRosterColumnHint(
            expected_header="Video",
            destination_team_key="main.cm.digital.video",
        ),
    )
)


@dataclass(frozen=True, slots=True)
class SvcaBethany2026ObservedTeamRosterColumn:
    observed_column: ObservedTeamRosterColumn
    scope: TeamRosterColumnScope
    matched_hint: TeamRosterColumnHint | None


@dataclass(frozen=True, slots=True)
class SvcaBethany2026TeamRosterColumnInventory:
    integration_key: str
    source_contract_revision: str
    filename: str
    workbook_sha256: str
    sheet_name: str
    observed_columns: tuple[SvcaBethany2026ObservedTeamRosterColumn, ...]
    parsed_workbook: ParsedWorshipWorkbook
    target_matches: tuple[TargetMatch, ...]


def _inventory_headers(content):
    workbook = load_workbook(
        BytesIO(content), data_only=False, read_only=True, keep_links=False
    )
    try:
        sheet = workbook[SUPPORTED_SHEET]
        observed = []
        for column_number in range(1, TEAM_ROSTER_INVENTORY_LAST_COLUMN + 1):
            column = get_column_letter(column_number)
            raw_header = sheet.cell(TEAM_ROSTER_HEADER_ROW, column_number).value
            observed_header = "" if raw_header is None else raw_header
            observed_column = ObservedTeamRosterColumn(
                sheet_name=sheet.title,
                column=column,
                observed_header=observed_header,
            )
            scope = TEAM_ROSTER_COLUMN_SCOPE[column]
            observed.append(
                SvcaBethany2026ObservedTeamRosterColumn(
                    observed_column=observed_column,
                    scope=scope,
                    matched_hint=(
                        resolve_exact_team_roster_column_hint(
                            observed_header, TEAM_ROSTER_COLUMN_HINTS
                        )
                        if scope
                        == TeamRosterColumnScope.TARGET_PROFILE_REVIEW_CANDIDATE
                        else None
                    ),
                )
            )
        return tuple(observed)
    finally:
        workbook.close()


def inventory_team_roster_columns(content, *, filename="workbook.xlsx"):
    """Return zero-write column and exact target evidence for this adapter."""

    parsed = parse_known_worship_workbook(content, filename=filename)
    observed_columns = _inventory_headers(content)
    target_matches = match_exact_service_event_targets(parsed)
    return SvcaBethany2026TeamRosterColumnInventory(
        integration_key=INTEGRATION_KEY,
        source_contract_revision=CONTRACT_REVISION,
        filename=parsed.filename,
        workbook_sha256=parsed.sha256,
        sheet_name=SUPPORTED_SHEET,
        observed_columns=observed_columns,
        parsed_workbook=parsed,
        target_matches=target_matches,
    )
