"""Zero-write Sound assignment identity review for the known annual workbook.

MO-S.6F.1A consumes the existing strict Worship workbook parser and target
matcher without changing their contracts.  It adds only the named adapter's
Column F source semantic, exact-team TeamMembership review, current assignment
classification, and expiring signed preview evidence.  There is deliberately
no confirmation or write operation in this module.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from enum import StrEnum
from hashlib import sha256
from io import BytesIO
import re
import unicodedata

from django.core import signing
from django.contrib.auth import get_user_model
from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import load_workbook

from accounts.structure_selectors import user_matches_structure_audience
from events.models import ServiceEvent, service_event_is_history
from events.service_profile_readiness import service_event_audience_readiness
from events.service_profile_runtime import inspect_service_profile_identity

from ..models import (
    MinistryTeam,
    TeamAssignment,
    TeamAssignmentMember,
    TeamMembership,
)
from .worship_xlsx_preview import (
    CONTRACT_REVISION as WORSHIP_PARSER_CONTRACT_REVISION,
    INTEGRATION_KEY,
    SIGNING_MAX_AGE_SECONDS,
    SUPPORTED_EVENT_TYPE,
    SUPPORTED_LOCAL_TIME,
    SUPPORTED_PROFILE_KEY,
    SUPPORTED_ROWS,
    SUPPORTED_SHEET,
    TOKEN_ORDER,
    ParsedWorshipWorkbook,
    SignedWorkbookStateError,
    TargetServiceProfileError,
    TargetMatchState,
    WorkbookContractError,
    WorkbookErrorCode,
    match_exact_service_event_targets,
    parse_known_worship_workbook,
    resolve_target_service_profile,
)
from .worship_context_review import FINGERPRINT_RE as WORSHIP_CONTEXT_FINGERPRINT_RE


SOURCE_CONTRACT_REVISION = "SVCA_BETHANY_0930_2026_SOUND_SOURCE_V1"
MAPPING_STATE_CONTRACT_REVISION = "SVCA_BETHANY_0930_2026_SOUND_MAPPING_V1"
PREVIEW_CONTRACT_REVISION = "SVCA_BETHANY_0930_2026_SOUND_PREVIEW_V1"
MAPPING_SIGNING_VERSION = 1
PREVIEW_SIGNING_VERSION = 1
MAPPING_SIGNING_SALT = "ministry.sound-assignment-mapping.v1"
PREVIEW_SIGNING_SALT = "ministry.sound-assignment-preview.v1"
SOURCE_SEMANTIC = "sound"
SOURCE_COLUMN = "F"
EXPECTED_SOUND_HEADER = "Sound"

# Named-adapter/deployment fact only.  Generic resolvers below accept a key and
# never branch on this literal or infer a team from its display name/taxonomy.
SVCA_SOUND_TEAM_KEY = "main.cm.digital.sound"

CURRENT_ASSIGNMENT_STATUSES = (
    TeamAssignment.STATUS_SCHEDULED,
    TeamAssignment.STATUS_CONFIRMED,
    TeamAssignment.STATUS_PREPARED,
)
HISTORICAL_ASSIGNMENT_STATUSES = (
    TeamAssignment.STATUS_COMPLETED,
    TeamAssignment.STATUS_CANCELLED,
)

_SHA256_RE = re.compile(r"^[0-9A-F]{64}$")
_PLACEHOLDER_RE = re.compile(
    r"(?:^|\s)(?:tbd|tba|to\s+be\s+(?:decided|assigned|confirmed))(?:$|\s)",
    re.IGNORECASE,
)
_REPLACEMENT_RE = re.compile(
    r"(?:^|\s)(?:sub|substitute|replacement|replace|instead)(?:$|\s)",
    re.IGNORECASE,
)
_UNSUPPORTED_SEPARATORS = ("/", "&", ",", "，", "、", ";", "；")
_ANNOTATION_MARKERS = ("(", ")", "（", "）", "[", "]", "【", "】", ":", "：")
_REPLACEMENT_MARKERS = ("->", "=>", "→", "替补", "替補", "代班", "待定", "待确认", "待確認")


class SoundSourceState(StrEnum):
    NO_SOURCE_PROPOSAL = "no_source_proposal"
    SUPPORTED_LITERAL = "supported_literal"
    FORMULA_BLOCKED = "formula_blocked"
    ERROR_BLOCKED = "error_blocked"
    UNSUPPORTED_TOKEN = "unsupported_token"


class SoundIdentityState(StrEnum):
    NOT_APPLICABLE = "not_applicable"
    EXACT_PREFILL_AVAILABLE = "exact_prefill_available"
    MAPPING_REQUIRED = "mapping_required"
    AMBIGUOUS = "ambiguous"
    REVIEWED_SELECTED = "reviewed_selected"


class SoundTargetState(StrEnum):
    NO_SOURCE_PROPOSAL = "no_source_proposal"
    CREATE_CANDIDATE = "create_candidate"
    FILL_CANDIDATE = "fill_candidate"
    REPLACE_CANDIDATE = "replace_candidate"
    EXACT_NOOP = "exact_noop"
    EXISTING_ROSTER_BLOCKER = "existing_roster_blocker"
    DUPLICATE_ASSIGNMENT_BLOCKER = "duplicate_assignment_blocker"
    HISTORICAL_ASSIGNMENT_BLOCKER = "historical_assignment_blocker"
    UNKNOWN_ASSIGNMENT_BLOCKER = "unknown_assignment_blocker"
    INVALID_ASSIGNMENT_BLOCKER = "invalid_assignment_blocker"
    HISTORICAL_EVENT_BLOCKER = "historical_event_blocker"
    INVALID_TARGET_BLOCKER = "invalid_target_blocker"
    SOURCE_BLOCKER = "source_blocker"
    IDENTITY_BLOCKER = "identity_blocker"
    AUDIENCE_SAFETY_BLOCKER = "audience_safety_blocker"


class SoundDestinationTeamErrorCode(StrEnum):
    MISSING = "sound_team_missing"
    DUPLICATE = "sound_team_duplicate"
    INACTIVE = "sound_team_inactive"
    NON_ASSIGNABLE = "sound_team_non_assignable"
    IDENTITY_CHANGED = "sound_team_identity_changed"


class SoundDestinationTeamError(RuntimeError):
    def __init__(self, code):
        self.code = code
        super().__init__(f"Configured Sound team is unavailable: {code.value}.")


class SoundMappingStateError(ValueError):
    pass


class SoundMappingValidationError(ValueError):
    pass


class SignedSoundPreviewError(ValueError):
    pass


@dataclass(frozen=True)
class SoundSourceRow:
    source_row: int
    source_cell: str
    local_date: date
    date_kind: str
    worship_token: str
    source_state: SoundSourceState
    normalized_token: str | None


@dataclass(frozen=True)
class ParsedSoundAssignmentWorkbook:
    worship: ParsedWorshipWorkbook
    rows: tuple[SoundSourceRow, ...]

    @property
    def filename(self):
        return self.worship.filename

    @property
    def sha256(self):
        return self.worship.sha256


@dataclass(frozen=True)
class SoundMembershipCandidate:
    membership_id: int
    visible_identity: str
    identity_digest: str
    linked_user_id: int | None
    linked_state: str
    linked_user_active: bool | None
    updated_at: str


@dataclass(frozen=True)
class SoundTokenReview:
    index: int
    token: str
    occurrence_count: int
    exact_match_ids: tuple[int, ...]
    identity_state: SoundIdentityState
    prefill_membership_id: int | None


@dataclass(frozen=True)
class SoundMappingReview:
    parsed: ParsedSoundAssignmentWorkbook
    team: MinistryTeam
    candidates: tuple[SoundMembershipCandidate, ...]
    token_reviews: tuple[SoundTokenReview, ...]
    signed_state: str


@dataclass(frozen=True)
class SoundAssignmentPreviewRow:
    source: SoundSourceRow
    identity_state: SoundIdentityState
    selected_membership: SoundMembershipCandidate | None
    event: ServiceEvent | None
    target_state: SoundTargetState
    blocker_detail: str | None
    current_assignment_ids: tuple[int, ...]
    current_roster_membership_ids: tuple[int, ...]
    historical_assignment_ids: tuple[int, ...]
    current_membership: SoundMembershipCandidate | None


@dataclass(frozen=True)
class SoundAssignmentPreview:
    mapping_review: SoundMappingReview
    rows: tuple[SoundAssignmentPreviewRow, ...]
    normalized_payload: dict
    signed_payload: str

    @property
    def signed_payload_bytes(self):
        return len(self.signed_payload.encode("utf-8"))

    @property
    def create_candidate_count(self):
        return sum(row.target_state == SoundTargetState.CREATE_CANDIDATE for row in self.rows)

    @property
    def fill_candidate_count(self):
        return sum(row.target_state == SoundTargetState.FILL_CANDIDATE for row in self.rows)

    @property
    def replace_candidate_count(self):
        return sum(
            row.target_state == SoundTargetState.REPLACE_CANDIDATE for row in self.rows
        )

    @property
    def exact_noop_count(self):
        return sum(row.target_state == SoundTargetState.EXACT_NOOP for row in self.rows)

    @property
    def no_source_count(self):
        return sum(row.target_state == SoundTargetState.NO_SOURCE_PROPOSAL for row in self.rows)

    @property
    def historical_event_count(self):
        return sum(
            row.target_state == SoundTargetState.HISTORICAL_EVENT_BLOCKER
            for row in self.rows
        )

    @property
    def hard_blocker_count(self):
        return sum(
            row.target_state
            not in {
                SoundTargetState.CREATE_CANDIDATE,
                SoundTargetState.FILL_CANDIDATE,
                SoundTargetState.REPLACE_CANDIDATE,
                SoundTargetState.EXACT_NOOP,
                SoundTargetState.NO_SOURCE_PROPOSAL,
                SoundTargetState.HISTORICAL_EVENT_BLOCKER,
            }
            for row in self.rows
        )

    @property
    def blocked_count(self):
        """Compatibility name for confirmation-suppressing hard blockers."""

        return self.hard_blocker_count

    @property
    def is_confirmable(self):
        return (
            self.create_candidate_count
            + self.fill_candidate_count
            + self.replace_candidate_count
            > 0
            and self.hard_blocker_count == 0
        )


def user_can_preview_sound_assignments(user):
    return bool(
        getattr(user, "is_authenticated", False)
        and getattr(user, "is_active", False)
        and (
            getattr(user, "is_staff", False)
            or getattr(user, "is_superuser", False)
        )
    )


def normalize_sound_identity(value):
    """Column-F identity normalization: Unicode NFC plus outer trim only."""

    if not isinstance(value, str):
        raise TypeError("Sound identity must be text.")
    return unicodedata.normalize("NFC", value).strip()


def _identity_digest(value):
    return sha256(normalize_sound_identity(value).encode("utf-8")).hexdigest().upper()


def _token_has_unsupported_syntax(token):
    return bool(
        "\n" in token
        or "\r" in token
        or "\t" in token
        or len(token) > 120
        or any(marker in token for marker in _UNSUPPORTED_SEPARATORS)
        or any(marker in token for marker in _ANNOTATION_MARKERS)
        or any(marker in token for marker in _REPLACEMENT_MARKERS)
        or _PLACEHOLDER_RE.search(token)
        or _REPLACEMENT_RE.search(token)
    )


def _classify_sound_cell(cell):
    if cell.data_type == "f":
        return SoundSourceState.FORMULA_BLOCKED, None
    if cell.data_type == "e":
        return SoundSourceState.ERROR_BLOCKED, None
    if cell.value is None:
        return SoundSourceState.NO_SOURCE_PROPOSAL, None
    if not isinstance(cell.value, str):
        return SoundSourceState.UNSUPPORTED_TOKEN, None

    token = normalize_sound_identity(cell.value)
    if not token:
        return SoundSourceState.NO_SOURCE_PROPOSAL, None
    if _token_has_unsupported_syntax(token):
        return SoundSourceState.UNSUPPORTED_TOKEN, token
    return SoundSourceState.SUPPORTED_LITERAL, token


def parse_known_sound_assignment_workbook(content, *, filename="workbook.xlsx"):
    """Parse Column F after the unchanged strict annual-workbook parser passes."""

    worship = parse_known_worship_workbook(content, filename=filename)
    try:
        source = load_workbook(
            BytesIO(content), data_only=False, read_only=False, keep_links=False
        )
    except Exception as exc:
        raise WorkbookContractError(
            WorkbookErrorCode.INVALID_XLSX,
            "The uploaded file is not a readable XLSX workbook.",
        ) from exc
    try:
        sheet = source[SUPPORTED_SHEET]
        if sheet[f"{SOURCE_COLUMN}3"].value != EXPECTED_SOUND_HEADER:
            raise WorkbookContractError(
                WorkbookErrorCode.HEADER_MISMATCH,
                "Cell F3 does not match the supported Sound assignment header.",
            )
        rows = []
        for worship_row in worship.rows:
            cell = sheet[f"{SOURCE_COLUMN}{worship_row.source_row}"]
            source_state, token = _classify_sound_cell(cell)
            rows.append(
                SoundSourceRow(
                    source_row=worship_row.source_row,
                    source_cell=cell.coordinate,
                    local_date=worship_row.local_date,
                    date_kind=worship_row.date_kind,
                    worship_token=worship_row.token,
                    source_state=source_state,
                    normalized_token=token,
                )
            )
        return ParsedSoundAssignmentWorkbook(worship=worship, rows=tuple(rows))
    finally:
        source.close()


def resolve_destination_team(team_key):
    """Resolve one exact active assignable team by stable key, with no fallback."""

    try:
        team = MinistryTeam.objects.get(team_key=team_key)
    except MinistryTeam.DoesNotExist as exc:
        raise SoundDestinationTeamError(SoundDestinationTeamErrorCode.MISSING) from exc
    except MinistryTeam.MultipleObjectsReturned as exc:
        raise SoundDestinationTeamError(SoundDestinationTeamErrorCode.DUPLICATE) from exc
    if not team.is_active:
        raise SoundDestinationTeamError(SoundDestinationTeamErrorCode.INACTIVE)
    if not team.is_assignable:
        raise SoundDestinationTeamError(SoundDestinationTeamErrorCode.NON_ASSIGNABLE)
    return team


def _privacy_safe_visible_identity(membership):
    """Canonical visible serving identity without email/contact fallbacks."""

    if membership.display_name:
        value = membership.display_name
    elif membership.user_id:
        value = membership.user.get_full_name() or membership.user.get_username()
    else:
        raise SoundMappingStateError(
            "An active Sound membership has no privacy-safe visible identity."
        )
    value = normalize_sound_identity(value)
    if not value:
        raise SoundMappingStateError(
            "An active Sound membership has no privacy-safe visible identity."
        )
    return value


def _membership_candidate(membership):
    visible_identity = _privacy_safe_visible_identity(membership)
    return SoundMembershipCandidate(
        membership_id=membership.pk,
        visible_identity=visible_identity,
        identity_digest=_identity_digest(visible_identity),
        linked_user_id=membership.user_id,
        linked_state="linked" if membership.user_id else "display_name_only",
        linked_user_active=(membership.user.is_active if membership.user_id else None),
        updated_at=membership.updated_at.isoformat(),
    )


def _current_candidates(team):
    memberships = TeamMembership.objects.filter(
        team=team, is_active=True
    ).select_related("user").order_by("id")
    return tuple(_membership_candidate(item) for item in memberships)


def _token_reviews(parsed, candidates):
    counts = Counter(
        row.normalized_token
        for row in parsed.rows
        if row.source_state == SoundSourceState.SUPPORTED_LITERAL
    )
    first_seen = []
    for row in parsed.rows:
        if (
            row.source_state == SoundSourceState.SUPPORTED_LITERAL
            and row.normalized_token not in first_seen
        ):
            first_seen.append(row.normalized_token)
    reviews = []
    for index, token in enumerate(first_seen):
        exact_ids = tuple(
            candidate.membership_id
            for candidate in candidates
            if normalize_sound_identity(candidate.visible_identity) == token
        )
        if len(exact_ids) == 1:
            state = SoundIdentityState.EXACT_PREFILL_AVAILABLE
            prefill = exact_ids[0]
        elif len(exact_ids) > 1:
            state = SoundIdentityState.AMBIGUOUS
            prefill = None
        else:
            state = SoundIdentityState.MAPPING_REQUIRED
            prefill = None
        reviews.append(
            SoundTokenReview(
                index=index,
                token=token,
                occurrence_count=counts[token],
                exact_match_ids=exact_ids,
                identity_state=state,
                prefill_membership_id=prefill,
            )
        )
    return tuple(reviews)


def _source_row_payload(row):
    return {
        "source_row": row.source_row,
        "source_cell": row.source_cell,
        "local_date": row.local_date.isoformat(),
        "date_kind": row.date_kind,
        "worship_token": row.worship_token,
        "source_state": row.source_state.value,
        "normalized_token": row.normalized_token,
    }


def _candidate_payload(candidate):
    return {
        "membership_id": candidate.membership_id,
        "visible_identity": candidate.visible_identity,
        "identity_digest": candidate.identity_digest,
        "linked_user_id": candidate.linked_user_id,
        "linked_state": candidate.linked_state,
        "linked_user_active": candidate.linked_user_active,
        "updated_at": candidate.updated_at,
    }


def prepare_sound_assignment_mapping(*, content, filename, user):
    parsed = parse_known_sound_assignment_workbook(content, filename=filename)
    profile = resolve_target_service_profile()
    team = resolve_destination_team(SVCA_SOUND_TEAM_KEY)
    candidates = _current_candidates(team)
    reviews = _token_reviews(parsed, candidates)
    payload = {
        "contract_revision": MAPPING_STATE_CONTRACT_REVISION,
        "signing_version": MAPPING_SIGNING_VERSION,
        "state_type": "sound_assignment_mapping",
        "source_contract_revision": SOURCE_CONTRACT_REVISION,
        "worship_parser_contract_revision": WORSHIP_PARSER_CONTRACT_REVISION,
        "integration_key": INTEGRATION_KEY,
        "source_semantic": SOURCE_SEMANTIC,
        "source_column": SOURCE_COLUMN,
        "generated_at": timezone.now().isoformat(),
        "user_id": user.pk,
        "filename": parsed.filename,
        "sha256": parsed.sha256,
        "profile": {
            "id": profile.pk,
            "key": profile.key,
            "event_type": profile.event_type,
        },
        "team": {
            "id": team.pk,
            "key": team.team_key,
            "active": team.is_active,
            "assignable": team.is_assignable,
        },
        "rows": [_source_row_payload(row) for row in parsed.rows],
        "candidates": [_candidate_payload(item) for item in candidates],
    }
    signed_state = signing.dumps(payload, compress=True, salt=MAPPING_SIGNING_SALT)
    return SoundMappingReview(
        parsed=parsed,
        team=team,
        candidates=candidates,
        token_reviews=reviews,
        signed_state=signed_state,
    )


def _safe_filename(value):
    return str(value or "workbook.xlsx").replace("\\", "/").rsplit("/", 1)[-1]


def _decode_mapping_payload(token, *, user, max_age):
    try:
        payload = signing.loads(token, salt=MAPPING_SIGNING_SALT, max_age=max_age)
    except signing.BadSignature as exc:
        raise SoundMappingStateError(
            "Sound mapping state is invalid or expired. Upload the workbook again."
        ) from exc
    required = {
        "contract_revision",
        "signing_version",
        "state_type",
        "source_contract_revision",
        "worship_parser_contract_revision",
        "integration_key",
        "source_semantic",
        "source_column",
        "generated_at",
        "user_id",
        "filename",
        "sha256",
        "profile",
        "team",
        "rows",
        "candidates",
    }
    if (
        not isinstance(payload, dict)
        or set(payload) != required
        or payload["contract_revision"] != MAPPING_STATE_CONTRACT_REVISION
        or payload["signing_version"] != MAPPING_SIGNING_VERSION
        or payload["state_type"] != "sound_assignment_mapping"
        or payload["source_contract_revision"] != SOURCE_CONTRACT_REVISION
        or payload["worship_parser_contract_revision"]
        != WORSHIP_PARSER_CONTRACT_REVISION
        or payload["integration_key"] != INTEGRATION_KEY
        or payload["source_semantic"] != SOURCE_SEMANTIC
        or payload["source_column"] != SOURCE_COLUMN
        or payload["user_id"] != getattr(user, "pk", None)
        or not isinstance(payload["generated_at"], str)
        or not isinstance(payload["filename"], str)
        or _safe_filename(payload["filename"]) != payload["filename"]
        or not isinstance(payload["sha256"], str)
        or _SHA256_RE.fullmatch(payload["sha256"]) is None
        or not isinstance(payload["rows"], list)
        or len(payload["rows"]) != len(SUPPORTED_ROWS)
        or not isinstance(payload["candidates"], list)
    ):
        raise SoundMappingStateError("Sound mapping state is malformed.")
    return payload


def decode_sound_assignment_mapping(
    token, *, user, max_age=SIGNING_MAX_AGE_SECONDS
):
    payload = _decode_mapping_payload(token, user=user, max_age=max_age)
    profile = resolve_target_service_profile()
    if payload["profile"] != {
        "id": profile.pk,
        "key": profile.key,
        "event_type": profile.event_type,
    }:
        raise SoundMappingStateError("The target Service Profile changed.")
    team = resolve_destination_team(SVCA_SOUND_TEAM_KEY)
    if payload["team"] != {
        "id": team.pk,
        "key": team.team_key,
        "active": team.is_active,
        "assignable": team.is_assignable,
    }:
        raise SoundMappingStateError("The configured Sound team changed.")
    current_candidates = _current_candidates(team)
    if payload["candidates"] != [_candidate_payload(item) for item in current_candidates]:
        raise SoundMappingStateError(
            "Sound membership identity changed. Upload and review the workbook again."
        )

    allowed_row_keys = {
        "source_row",
        "source_cell",
        "local_date",
        "date_kind",
        "worship_token",
        "source_state",
        "normalized_token",
    }
    rows = []
    try:
        for item in payload["rows"]:
            if not isinstance(item, dict) or set(item) != allowed_row_keys:
                raise ValueError
            row = SoundSourceRow(
                source_row=item["source_row"],
                source_cell=item["source_cell"],
                local_date=date.fromisoformat(item["local_date"]),
                date_kind=item["date_kind"],
                worship_token=item["worship_token"],
                source_state=SoundSourceState(item["source_state"]),
                normalized_token=item["normalized_token"],
            )
            rows.append(row)
    except (KeyError, TypeError, ValueError) as exc:
        raise SoundMappingStateError("Sound mapping rows are malformed.") from exc

    if tuple(row.source_row for row in rows) != SUPPORTED_ROWS:
        raise SoundMappingStateError("Sound mapping rows are not canonical.")
    for row in rows:
        expected_cell = f"{SOURCE_COLUMN}{row.source_row}"
        index = 0 if row.source_row == 4 else row.source_row - 5
        expected_date = date(2026, 1, 4) + timedelta(weeks=index)
        if (
            type(row.source_row) is not int
            or row.source_cell != expected_cell
            or row.local_date != expected_date
            or row.date_kind
            != ("literal" if row.source_row == 4 else "formula_cached")
            or row.worship_token not in TOKEN_ORDER
            or row.source_state == SoundSourceState.SUPPORTED_LITERAL
            and (
                not row.normalized_token
                or _token_has_unsupported_syntax(row.normalized_token)
            )
            or row.normalized_token is not None
            and normalize_sound_identity(row.normalized_token) != row.normalized_token
            or row.source_state != SoundSourceState.SUPPORTED_LITERAL
            and row.source_state != SoundSourceState.UNSUPPORTED_TOKEN
            and row.normalized_token is not None
        ):
            raise SoundMappingStateError("Sound mapping rows are not canonical.")

    worship_rows = []
    from .worship_xlsx_preview import ParsedWorkbookRow

    for row in rows:
        worship_rows.append(
            ParsedWorkbookRow(
                source_row=row.source_row,
                source_cell=f"A{row.source_row}/B{row.source_row}",
                local_date=row.local_date,
                date_kind=row.date_kind,
                token=row.worship_token,
            )
        )
    worship = ParsedWorshipWorkbook(
        filename=payload["filename"],
        sha256=payload["sha256"],
        rows=tuple(worship_rows),
        token_counts=dict(Counter(row.worship_token for row in rows)),
        unsupported_rows=(
            {"row": 5, "kind": "friday", "date": "2026-01-09"},
            {"row": 57, "kind": "spillover", "date": "2027-01-03"},
            {"row": 58, "kind": "spillover", "date": "2027-01-10"},
        ),
    )
    parsed = ParsedSoundAssignmentWorkbook(worship=worship, rows=tuple(rows))
    return SoundMappingReview(
        parsed=parsed,
        team=team,
        candidates=current_candidates,
        token_reviews=_token_reviews(parsed, current_candidates),
        signed_state=token,
    )


def _validate_selected_mapping(review, selected_mapping):
    candidate_by_id = {item.membership_id: item for item in review.candidates}
    allowed_tokens = {item.token for item in review.token_reviews}
    if set(selected_mapping) - allowed_tokens:
        raise SoundMappingValidationError("Sound mapping contains an unknown token.")
    resolved = {}
    for token in allowed_tokens:
        raw_id = selected_mapping.get(token)
        if raw_id in (None, ""):
            resolved[token] = None
            continue
        try:
            membership_id = int(raw_id)
        except (TypeError, ValueError) as exc:
            raise SoundMappingValidationError("Sound membership ID is malformed.") from exc
        candidate = candidate_by_id.get(membership_id)
        if candidate is None:
            raise SoundMappingValidationError(
                "Selected membership is missing, inactive, or belongs to another team."
            )
        resolved[token] = candidate
    return resolved


def _assignment_queryset(event_ids, team):
    members = TeamAssignmentMember.objects.select_related(
        "membership", "membership__user"
    ).order_by("id")
    return (
        TeamAssignment.objects.filter(service_event_id__in=event_ids, ministry_team=team)
        .select_related("service_event", "ministry_team")
        .prefetch_related(Prefetch("assignment_members", queryset=members))
        .order_by("id")
    )


def _assignment_baseline(assignments):
    return [
        {
            "id": assignment.pk,
            "status": assignment.status,
            "updated_at": assignment.updated_at.isoformat(),
            "reviewed_worship_context_fingerprint": (
                assignment.reviewed_worship_context_fingerprint
            ),
            "members": [
                {
                    "assignment_member_id": item.pk,
                    "membership_id": item.membership_id,
                    "membership_active": item.membership.is_active,
                    "confirmed": item.confirmed_at is not None,
                }
                for item in assignment.assignment_members.all()
            ],
        }
        for assignment in assignments
    ]


def _classify_assignment(assignments, selected_membership_id):
    current = [item for item in assignments if item.status in CURRENT_ASSIGNMENT_STATUSES]
    historical = [
        item for item in assignments if item.status in HISTORICAL_ASSIGNMENT_STATUSES
    ]
    unknown = [
        item
        for item in assignments
        if item.status
        not in {*CURRENT_ASSIGNMENT_STATUSES, *HISTORICAL_ASSIGNMENT_STATUSES}
    ]
    current_ids = tuple(item.pk for item in current)
    historical_ids = tuple(item.pk for item in historical)
    roster = ()
    if len(current) == 1:
        roster = tuple(
            sorted(
                item.membership_id
                for item in current[0].assignment_members.all()
            )
        )
    if len(current) > 1:
        state = SoundTargetState.DUPLICATE_ASSIGNMENT_BLOCKER
    elif unknown:
        state = SoundTargetState.UNKNOWN_ASSIGNMENT_BLOCKER
    elif historical:
        state = SoundTargetState.HISTORICAL_ASSIGNMENT_BLOCKER
    elif len(current) == 1:
        assignment = current[0]
        members = list(assignment.assignment_members.all())
        fingerprint = assignment.reviewed_worship_context_fingerprint
        if (
            fingerprint is not None
            and WORSHIP_CONTEXT_FINGERPRINT_RE.fullmatch(fingerprint) is None
        ):
            state = SoundTargetState.INVALID_ASSIGNMENT_BLOCKER
        elif any(not item.membership.is_active for item in members):
            state = SoundTargetState.EXISTING_ROSTER_BLOCKER
        elif len(members) > 1:
            state = SoundTargetState.EXISTING_ROSTER_BLOCKER
        elif len(members) == 1 and members[0].membership_id == selected_membership_id:
            state = SoundTargetState.EXACT_NOOP
        elif assignment.status != TeamAssignment.STATUS_SCHEDULED:
            state = SoundTargetState.EXISTING_ROSTER_BLOCKER
        elif not members:
            state = SoundTargetState.FILL_CANDIDATE
        elif (
            members[0].confirmed_at is None
            and members[0].confirmation_note == ""
        ):
            state = SoundTargetState.REPLACE_CANDIDATE
        else:
            state = SoundTargetState.EXISTING_ROSTER_BLOCKER
    else:
        state = SoundTargetState.CREATE_CANDIDATE
    return state, current_ids, roster, historical_ids


def _target_detail(target_state):
    return {
        TargetMatchState.NO_TARGET: "The exact target event is missing.",
        TargetMatchState.MULTIPLE_EXACT_TARGETS: "Multiple exact target events exist.",
        TargetMatchState.TARGET_EVENT_PROFILE_FK_MISSING: "The target event has no canonical Service Profile FK.",
        TargetMatchState.TARGET_EVENT_PROFILE_IDENTITY_DRIFT: "The target event Service Profile identity drifted.",
        TargetMatchState.TARGET_EVENT_OWNED_BY_OTHER_PROFILE: "The target event belongs to another Service Profile.",
        TargetMatchState.LIFECYCLE_CONFLICT: "The target event lifecycle is unsupported.",
        TargetMatchState.AUDIENCE_INVALID_CONFLICT: "The target event audience is not ready.",
    }.get(target_state)


def _event_is_historical_for_assignment_import(event, *, now):
    """V1 never backfills completed or canonically elapsed ServiceEvents."""

    return (
        event.status == ServiceEvent.STATUS_COMPLETED
        or service_event_is_history(event, now=now)
    )


def _membership_is_outside_event_audience(candidate, event):
    """Mirror interactive audience matching, but never permit bulk override."""

    if not candidate.linked_user_id or not candidate.linked_user_active:
        return False
    audience_units = list(event.get_audience_scope_units())
    return bool(audience_units) and not user_matches_structure_audience(
        get_user_model()._default_manager.get(pk=candidate.linked_user_id),
        audience_units,
    )


def _event_payload(event):
    local_start = timezone.localtime(event.start_datetime)
    return {
        "id": event.pk,
        "service_profile_id": event.service_profile_id,
        "expected_profile_key": event.service_profile.key,
        "scheduling_revision": event.scheduling_revision,
        "local_date": local_start.date().isoformat(),
        "local_time": local_start.time().replace(tzinfo=None).isoformat(),
        "event_type": event.event_type,
        "status": event.status,
    }


def build_sound_assignment_preview(*, mapping_review, selected_mapping, user, now=None):
    if not user_can_preview_sound_assignments(user):
        raise SoundMappingValidationError("Sound assignment preview is staff-only.")
    team = resolve_destination_team(SVCA_SOUND_TEAM_KEY)
    if team.pk != mapping_review.team.pk:
        raise SoundMappingValidationError("The configured Sound team changed.")
    if mapping_review.candidates != _current_candidates(team):
        raise SoundMappingValidationError(
            "Sound membership identity changed. Upload and review the workbook again."
        )
    resolved_mapping = _validate_selected_mapping(mapping_review, selected_mapping)
    classification_now = now or timezone.now()
    matches = match_exact_service_event_targets(mapping_review.parsed.worship)
    event_ids = [match.event.pk for match in matches if match.event is not None]
    assignments_by_event = defaultdict(list)
    for assignment in _assignment_queryset(event_ids, team):
        assignments_by_event[assignment.service_event_id].append(assignment)

    token_review_by_token = {item.token: item for item in mapping_review.token_reviews}
    rows = []
    proposal_rows = []
    for source, match in zip(mapping_review.parsed.rows, matches, strict=True):
        selected = (
            resolved_mapping.get(source.normalized_token)
            if source.normalized_token is not None
            else None
        )
        event = match.event
        current_ids = ()
        roster = ()
        historical_ids = ()
        current_membership = None
        blocker_detail = None

        if source.source_state == SoundSourceState.NO_SOURCE_PROPOSAL:
            identity_state = SoundIdentityState.NOT_APPLICABLE
            target_state = SoundTargetState.NO_SOURCE_PROPOSAL
        elif source.source_state != SoundSourceState.SUPPORTED_LITERAL:
            identity_state = SoundIdentityState.NOT_APPLICABLE
            target_state = SoundTargetState.SOURCE_BLOCKER
            blocker_detail = "Column F is outside the one-person literal contract."
        elif selected is None:
            token_review = token_review_by_token[source.normalized_token]
            identity_state = token_review.identity_state
            target_state = SoundTargetState.IDENTITY_BLOCKER
            blocker_detail = "Explicit active Sound-team membership review is required."
        elif match.state != TargetMatchState.EXACT_TARGET_MATCHED:
            identity_state = SoundIdentityState.REVIEWED_SELECTED
            target_state = SoundTargetState.INVALID_TARGET_BLOCKER
            blocker_detail = _target_detail(match.state)
        else:
            identity_state = SoundIdentityState.REVIEWED_SELECTED
            assignments = assignments_by_event[event.pk]
            (
                assignment_state,
                current_ids,
                roster,
                historical_ids,
            ) = _classify_assignment(assignments, selected.membership_id)
            current = [
                item
                for item in assignments
                if item.status in CURRENT_ASSIGNMENT_STATUSES
            ]
            if len(current) == 1:
                members = list(current[0].assignment_members.all())
                if len(members) == 1:
                    current_membership = _membership_candidate(
                        members[0].membership
                    )
            if _event_is_historical_for_assignment_import(
                event, now=classification_now
            ):
                target_state = SoundTargetState.HISTORICAL_EVENT_BLOCKER
                blocker_detail = (
                    "The ServiceEvent is historical; V1 does not backfill assignments."
                )
            elif _membership_is_outside_event_audience(selected, event):
                target_state = SoundTargetState.AUDIENCE_SAFETY_BLOCKER
                blocker_detail = (
                    "The linked user is outside the event audience. Use the normal "
                    "manual assignment workflow for an intentional override."
                )
            else:
                target_state = assignment_state

            proposal_rows.append(
                {
                    "source": {
                        "row": source.source_row,
                        "cell": source.source_cell,
                        "semantic": SOURCE_SEMANTIC,
                        "literal_classification": source.source_state.value,
                        "token_digest": _identity_digest(source.normalized_token),
                    },
                    "event": _event_payload(event),
                    "team": {
                        "id": team.pk,
                        "key": team.team_key,
                        "active": team.is_active,
                        "assignable": team.is_assignable,
                    },
                    "membership": {
                        "id": selected.membership_id,
                        "team_id": team.pk,
                        "active": True,
                        "updated_at": selected.updated_at,
                        "linked_user_id": selected.linked_user_id,
                        "linked_user_active": selected.linked_user_active,
                        "visible_identity_digest": selected.identity_digest,
                    },
                    "target_state": target_state.value,
                    "assignment_baseline": _assignment_baseline(assignments),
                }
            )

        rows.append(
            SoundAssignmentPreviewRow(
                source=source,
                identity_state=identity_state,
                selected_membership=selected,
                event=event,
                target_state=target_state,
                blocker_detail=blocker_detail,
                current_assignment_ids=current_ids,
                current_roster_membership_ids=roster,
                historical_assignment_ids=historical_ids,
                current_membership=current_membership,
            )
        )

    profile = resolve_target_service_profile()
    payload = {
        "contract_revision": PREVIEW_CONTRACT_REVISION,
        "signing_version": PREVIEW_SIGNING_VERSION,
        "state_type": "sound_assignment_preview",
        "source_contract_revision": SOURCE_CONTRACT_REVISION,
        "worship_parser_contract_revision": WORSHIP_PARSER_CONTRACT_REVISION,
        "integration_key": INTEGRATION_KEY,
        "source_semantic": SOURCE_SEMANTIC,
        "source_column": SOURCE_COLUMN,
        "generated_at": timezone.now().isoformat(),
        "user_id": user.pk,
        "filename": mapping_review.parsed.filename,
        "sha256": mapping_review.parsed.sha256,
        "profile": {
            "id": profile.pk,
            "key": profile.key,
            "event_type": profile.event_type,
        },
        "team": {"id": team.pk, "key": team.team_key},
        "rows": proposal_rows,
    }
    signed_payload = signing.dumps(payload, compress=True, salt=PREVIEW_SIGNING_SALT)
    return SoundAssignmentPreview(
        mapping_review=mapping_review,
        rows=tuple(rows),
        normalized_payload=payload,
        signed_payload=signed_payload,
    )


def decode_signed_sound_assignment_preview(
    token, *, user, max_age=SIGNING_MAX_AGE_SECONDS, now=None
):
    """Strictly decode and revalidate the 1A evidence; performs no write."""

    try:
        payload = signing.loads(token, salt=PREVIEW_SIGNING_SALT, max_age=max_age)
    except signing.BadSignature as exc:
        raise SignedSoundPreviewError(
            "Signed Sound assignment preview is invalid or expired."
        ) from exc
    required = {
        "contract_revision",
        "signing_version",
        "state_type",
        "source_contract_revision",
        "worship_parser_contract_revision",
        "integration_key",
        "source_semantic",
        "source_column",
        "generated_at",
        "user_id",
        "filename",
        "sha256",
        "profile",
        "team",
        "rows",
    }
    if (
        not isinstance(payload, dict)
        or set(payload) != required
        or payload["contract_revision"] != PREVIEW_CONTRACT_REVISION
        or payload["signing_version"] != PREVIEW_SIGNING_VERSION
        or payload["state_type"] != "sound_assignment_preview"
        or payload["source_contract_revision"] != SOURCE_CONTRACT_REVISION
        or payload["worship_parser_contract_revision"]
        != WORSHIP_PARSER_CONTRACT_REVISION
        or payload["integration_key"] != INTEGRATION_KEY
        or payload["source_semantic"] != SOURCE_SEMANTIC
        or payload["source_column"] != SOURCE_COLUMN
        or payload["user_id"] != getattr(user, "pk", None)
        or not isinstance(payload["generated_at"], str)
        or not isinstance(payload["filename"], str)
        or _safe_filename(payload["filename"]) != payload["filename"]
        or not isinstance(payload["sha256"], str)
        or _SHA256_RE.fullmatch(payload["sha256"]) is None
        or not isinstance(payload["rows"], list)
    ):
        raise SignedSoundPreviewError("Signed Sound assignment preview is malformed.")

    profile = resolve_target_service_profile()
    if payload["profile"] != {
        "id": profile.pk,
        "key": profile.key,
        "event_type": profile.event_type,
    }:
        raise SignedSoundPreviewError("The target Service Profile changed.")
    team = resolve_destination_team(SVCA_SOUND_TEAM_KEY)
    if payload["team"] != {"id": team.pk, "key": team.team_key}:
        raise SignedSoundPreviewError("The configured Sound team changed.")

    classification_now = now or timezone.now()
    seen_cells = set()
    for row in payload["rows"]:
        try:
            source = row["source"]
            event_state = row["event"]
            membership_state = row["membership"]
            if (
                set(row)
                != {
                    "source",
                    "event",
                    "team",
                    "membership",
                    "target_state",
                    "assignment_baseline",
                }
                or source["semantic"] != SOURCE_SEMANTIC
                or source["literal_classification"]
                != SoundSourceState.SUPPORTED_LITERAL.value
                or source["cell"] != f"{SOURCE_COLUMN}{source['row']}"
                or source["row"] not in SUPPORTED_ROWS
                or source["cell"] in seen_cells
                or not isinstance(source["token_digest"], str)
                or _SHA256_RE.fullmatch(source["token_digest"]) is None
                or row["team"]
                != {
                    "id": team.pk,
                    "key": team.team_key,
                    "active": True,
                    "assignable": True,
                }
            ):
                raise ValueError
            seen_cells.add(source["cell"])
            event = ServiceEvent.objects.select_related("service_profile").get(
                pk=event_state["id"]
            )
            local_start = timezone.localtime(event.start_datetime)
            identity = inspect_service_profile_identity(event)
            if (
                event_state != _event_payload(event)
                or not identity.is_exact
                or event.service_profile_id != profile.pk
                or event.service_profile.key != SUPPORTED_PROFILE_KEY
                or event.scheduling_revision != event_state["scheduling_revision"]
                or event.event_type != SUPPORTED_EVENT_TYPE
                or local_start.date().isoformat() != event_state["local_date"]
                or local_start.time().replace(tzinfo=None).isoformat()
                != event_state["local_time"]
                or local_start.time().replace(tzinfo=None) != SUPPORTED_LOCAL_TIME
                or event.status != event_state["status"]
                or event.status
                not in {ServiceEvent.STATUS_PUBLISHED, ServiceEvent.STATUS_COMPLETED}
                or not service_event_audience_readiness(event)["ready"]
            ):
                raise ValueError
            membership = TeamMembership.objects.select_related("user").get(
                pk=membership_state["id"]
            )
            candidate = _membership_candidate(membership)
            assignments = list(_assignment_queryset((event.pk,), team))
            target_state, _current_ids, _roster, _historical_ids = (
                _classify_assignment(assignments, candidate.membership_id)
            )
            if _event_is_historical_for_assignment_import(
                event, now=classification_now
            ):
                target_state = SoundTargetState.HISTORICAL_EVENT_BLOCKER
            elif _membership_is_outside_event_audience(candidate, event):
                target_state = SoundTargetState.AUDIENCE_SAFETY_BLOCKER
            if (
                not membership.is_active
                or membership.team_id != team.pk
                or membership_state
                != {
                    "id": candidate.membership_id,
                    "team_id": team.pk,
                    "active": True,
                    "updated_at": candidate.updated_at,
                    "linked_user_id": candidate.linked_user_id,
                    "linked_user_active": candidate.linked_user_active,
                    "visible_identity_digest": candidate.identity_digest,
                }
                or row["target_state"] != target_state.value
                or row["assignment_baseline"] != _assignment_baseline(assignments)
            ):
                raise ValueError
        except (
            KeyError,
            TypeError,
            ValueError,
            ServiceEvent.DoesNotExist,
            TeamMembership.DoesNotExist,
            SoundMappingStateError,
        ) as exc:
            raise SignedSoundPreviewError(
                "Signed Sound assignment preview is stale or malformed."
            ) from exc
    return payload
