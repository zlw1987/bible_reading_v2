"""Zero-write Projection assignment review for Column E of the annual workbook.

This adapter reuses the strict annual-workbook structure and exact event matcher,
but owns a separate source grammar, identity review, roster comparison, and
signed-preview contract.  It deliberately exposes no confirmation or writer.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from enum import StrEnum
from hashlib import sha256
from io import BytesIO
import re
from types import SimpleNamespace
import unicodedata

from django.contrib.auth import get_user_model
from django.core import signing
from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import load_workbook

from accounts.structure_selectors import user_matches_structure_audience
from events.models import ServiceEvent, service_event_is_history
from events.service_profile_readiness import service_event_audience_readiness
from events.service_profile_runtime import inspect_service_profile_identity

from ..models import MinistryTeam, TeamAssignment, TeamAssignmentMember, TeamMembership
from .worship_context_review import FINGERPRINT_RE as WORSHIP_CONTEXT_FINGERPRINT_RE
from .worship_xlsx_preview import (
    CONTRACT_REVISION as WORSHIP_PARSER_CONTRACT_REVISION,
    INTEGRATION_KEY,
    SIGNING_MAX_AGE_SECONDS,
    SUPPORTED_EVENT_TYPE,
    SUPPORTED_LOCAL_TIME,
    SUPPORTED_PROFILE_KEY,
    SUPPORTED_ROWS,
    SUPPORTED_SHEET,
    TOKEN_ORDER,
    ParsedWorshipWorkbook,
    TargetMatchState,
    WorkbookContractError,
    WorkbookErrorCode,
    match_exact_service_event_targets,
    parse_known_worship_workbook,
    resolve_target_service_profile,
)


SOURCE_CONTRACT_REVISION = "SVCA_BETHANY_0930_2026_PROJECTION_SOURCE_V1"
MAPPING_STATE_CONTRACT_REVISION = "SVCA_BETHANY_0930_2026_PROJECTION_MAPPING_V1"
PREVIEW_CONTRACT_REVISION = "SVCA_BETHANY_0930_2026_PROJECTION_PREVIEW_V1"
MAPPING_SIGNING_VERSION = 1
PREVIEW_SIGNING_VERSION = 1
MAPPING_SIGNING_SALT = "ministry.projection-assignment-mapping.v1"
PREVIEW_SIGNING_SALT = "ministry.projection-assignment-preview.v1"
SOURCE_SEMANTIC = "projection"
SOURCE_COLUMN = "E"
EXPECTED_PROJECTION_HEADER = "projector"
SVCA_PROJECTION_TEAM_KEY = "main.cm.digital.projection"

CURRENT_ASSIGNMENT_STATUSES = (
    TeamAssignment.STATUS_SCHEDULED,
    TeamAssignment.STATUS_CONFIRMED,
    TeamAssignment.STATUS_PREPARED,
)
HISTORICAL_ASSIGNMENT_STATUSES = (
    TeamAssignment.STATUS_COMPLETED,
    TeamAssignment.STATUS_CANCELLED,
)

_SHA256_RE = re.compile(r"^[0-9A-F]{64}$")
_PLACEHOLDER_RE = re.compile(
    r"\b(?:tbd|tba)\b|\bto\s+be\s+(?:decided|assigned|confirmed)\b",
    re.IGNORECASE,
)
_REPLACEMENT_RE = re.compile(
    r"(?:^|\s)(?:sub|substitute|replacement|replace|instead)(?:$|\s)",
    re.IGNORECASE,
)
_UNSUPPORTED_SEPARATORS = ("&", ",", "，", "、", ";", "；", "／", "\\")
_ANNOTATION_MARKERS = (
    "(", ")", "（", "）", "[", "]", "【", "】", ":", "：",
)
_REPLACEMENT_MARKERS = (
    "->",
    "=>",
    "<-",
    "→",
    "←",
    "⇒",
    "替补",
    "替補",
    "代班",
    "待定",
    "待确认",
    "待確認",
)
MAX_SOURCE_LITERAL_LENGTH = 240
MAX_PERSON_TOKEN_LENGTH = 120


class ProjectionSourceState(StrEnum):
    NO_SOURCE_PROPOSAL = "no_source_proposal"
    SUPPORTED_LITERAL = "supported_literal"
    FORMULA_BLOCKED = "formula_blocked"
    ERROR_BLOCKED = "error_blocked"
    UNSUPPORTED_TOKEN = "unsupported_token"


class ProjectionIdentityState(StrEnum):
    NOT_APPLICABLE = "not_applicable"
    EXACT_PREFILL_AVAILABLE = "exact_prefill_available"
    MAPPING_REQUIRED = "mapping_required"
    AMBIGUOUS = "ambiguous"
    REVIEWED_SELECTED = "reviewed_selected"


class ProjectionTargetState(StrEnum):
    NO_SOURCE_PROPOSAL = "no_source_proposal"
    CREATE_CANDIDATE = "create_candidate"
    EXACT_NOOP = "exact_noop"
    EXISTING_ROSTER_BLOCKER = "existing_roster_blocker"
    DUPLICATE_ASSIGNMENT_BLOCKER = "duplicate_assignment_blocker"
    HISTORICAL_ASSIGNMENT_BLOCKER = "historical_assignment_blocker"
    UNKNOWN_ASSIGNMENT_BLOCKER = "unknown_assignment_blocker"
    INVALID_ASSIGNMENT_BLOCKER = "invalid_assignment_blocker"
    HISTORICAL_EVENT_BLOCKER = "historical_event_blocker"
    INVALID_TARGET_BLOCKER = "invalid_target_blocker"
    SOURCE_BLOCKER = "source_blocker"
    IDENTITY_BLOCKER = "identity_blocker"
    AUDIENCE_SAFETY_BLOCKER = "audience_safety_blocker"


class ProjectionDestinationTeamErrorCode(StrEnum):
    MISSING = "projection_team_missing"
    DUPLICATE = "projection_team_duplicate"
    INACTIVE = "projection_team_inactive"
    NON_ASSIGNABLE = "projection_team_non_assignable"


class ProjectionDestinationTeamError(RuntimeError):
    def __init__(self, code):
        self.code = code
        super().__init__(f"Configured Projection team is unavailable: {code.value}.")


class ProjectionMappingStateError(ValueError):
    pass


class ProjectionMappingValidationError(ValueError):
    pass


class SignedProjectionPreviewError(ValueError):
    pass


@dataclass(frozen=True)
class ProjectionSourceRow:
    source_row: int
    source_cell: str
    local_date: date
    date_kind: str
    worship_token: str
    source_state: ProjectionSourceState
    original_literal: str | None
    normalized_tokens: tuple[str, ...]


@dataclass(frozen=True)
class ParsedProjectionAssignmentWorkbook:
    worship: ParsedWorshipWorkbook
    rows: tuple[ProjectionSourceRow, ...]

    @property
    def filename(self):
        return self.worship.filename

    @property
    def sha256(self):
        return self.worship.sha256


@dataclass(frozen=True)
class ProjectionMembershipCandidate:
    membership_id: int
    visible_identity: str
    identity_digest: str
    linked_user_id: int | None
    linked_state: str
    linked_user_active: bool | None
    updated_at: str


@dataclass(frozen=True)
class ProjectionTokenReview:
    index: int
    token: str
    occurrence_count: int
    exact_match_ids: tuple[int, ...]
    identity_state: ProjectionIdentityState
    prefill_membership_id: int | None


@dataclass(frozen=True)
class ProjectionMappingReview:
    parsed: ParsedProjectionAssignmentWorkbook
    team: MinistryTeam
    candidates: tuple[ProjectionMembershipCandidate, ...]
    token_reviews: tuple[ProjectionTokenReview, ...]
    signed_state: str


@dataclass(frozen=True)
class ProjectionAssignmentPreviewRow:
    source: ProjectionSourceRow
    identity_state: ProjectionIdentityState
    selected_memberships: tuple[ProjectionMembershipCandidate, ...]
    event: ServiceEvent | None
    target_state: ProjectionTargetState
    blocker_detail: str | None
    current_assignment_ids: tuple[int, ...]
    current_roster_membership_ids: tuple[int, ...]
    historical_assignment_ids: tuple[int, ...]
    current_memberships: tuple[ProjectionMembershipCandidate, ...]


@dataclass(frozen=True)
class ProjectionAssignmentPreview:
    mapping_review: ProjectionMappingReview
    rows: tuple[ProjectionAssignmentPreviewRow, ...]
    normalized_payload: dict
    signed_payload: str

    @property
    def create_candidate_count(self):
        return sum(row.target_state == ProjectionTargetState.CREATE_CANDIDATE for row in self.rows)

    @property
    def exact_noop_count(self):
        return sum(row.target_state == ProjectionTargetState.EXACT_NOOP for row in self.rows)

    @property
    def no_source_count(self):
        return sum(row.target_state == ProjectionTargetState.NO_SOURCE_PROPOSAL for row in self.rows)

    @property
    def historical_event_count(self):
        return sum(row.target_state == ProjectionTargetState.HISTORICAL_EVENT_BLOCKER for row in self.rows)

    @property
    def hard_blocker_count(self):
        safe = {
            ProjectionTargetState.CREATE_CANDIDATE,
            ProjectionTargetState.EXACT_NOOP,
            ProjectionTargetState.NO_SOURCE_PROPOSAL,
            ProjectionTargetState.HISTORICAL_EVENT_BLOCKER,
        }
        return sum(row.target_state not in safe for row in self.rows)


def user_can_preview_projection_assignments(user):
    return bool(
        getattr(user, "is_authenticated", False)
        and getattr(user, "is_active", False)
        and (getattr(user, "is_staff", False) or getattr(user, "is_superuser", False))
    )


def normalize_projection_identity(value):
    if not isinstance(value, str):
        raise TypeError("Projection identity must be text.")
    return unicodedata.normalize("NFC", value).strip()


def _identity_digest(value):
    return sha256(normalize_projection_identity(value).encode("utf-8")).hexdigest().upper()


def _token_has_unsupported_syntax(token):
    return bool(
        not token
        or "\n" in token
        or "\r" in token
        or "\t" in token
        or len(token) > MAX_PERSON_TOKEN_LENGTH
        or any(marker in token for marker in _UNSUPPORTED_SEPARATORS)
        or any(marker in token for marker in _ANNOTATION_MARKERS)
        or any(marker in token for marker in _REPLACEMENT_MARKERS)
        or _PLACEHOLDER_RE.search(token)
        or _REPLACEMENT_RE.search(token)
    )


def _classify_projection_cell(cell):
    if cell.data_type == "f":
        return ProjectionSourceState.FORMULA_BLOCKED, None, ()
    if cell.data_type == "e":
        return ProjectionSourceState.ERROR_BLOCKED, None, ()
    if cell.value is None:
        return ProjectionSourceState.NO_SOURCE_PROPOSAL, None, ()
    if not isinstance(cell.value, str):
        return ProjectionSourceState.UNSUPPORTED_TOKEN, None, ()

    original = cell.value
    if not normalize_projection_identity(original):
        return ProjectionSourceState.NO_SOURCE_PROPOSAL, original, ()
    if len(original) > MAX_SOURCE_LITERAL_LENGTH or any(c in original for c in "\n\r\t"):
        return ProjectionSourceState.UNSUPPORTED_TOKEN, original, ()
    if original.count("/") > 1:
        return ProjectionSourceState.UNSUPPORTED_TOKEN, original, ()

    tokens = tuple(normalize_projection_identity(part) for part in original.split("/"))
    if len(tokens) not in (1, 2) or any(_token_has_unsupported_syntax(token) for token in tokens):
        return ProjectionSourceState.UNSUPPORTED_TOKEN, original, tokens
    if len(set(tokens)) != len(tokens):
        return ProjectionSourceState.UNSUPPORTED_TOKEN, original, tokens
    return ProjectionSourceState.SUPPORTED_LITERAL, original, tokens


def parse_known_projection_assignment_workbook(content, *, filename="workbook.xlsx"):
    """Run the unchanged annual parser, then interpret only E/projector."""

    worship = parse_known_worship_workbook(content, filename=filename)
    try:
        source = load_workbook(BytesIO(content), data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise WorkbookContractError(
            WorkbookErrorCode.INVALID_XLSX,
            "The uploaded file is not a readable XLSX workbook.",
        ) from exc
    try:
        sheet = source[SUPPORTED_SHEET]
        if sheet[f"{SOURCE_COLUMN}3"].value != EXPECTED_PROJECTION_HEADER:
            raise WorkbookContractError(
                WorkbookErrorCode.HEADER_MISMATCH,
                "Cell E3 does not match the supported Projection assignment header.",
            )
        rows = []
        for worship_row in worship.rows:
            cell = sheet[f"{SOURCE_COLUMN}{worship_row.source_row}"]
            state, original, tokens = _classify_projection_cell(cell)
            rows.append(
                ProjectionSourceRow(
                    source_row=worship_row.source_row,
                    source_cell=cell.coordinate,
                    local_date=worship_row.local_date,
                    date_kind=worship_row.date_kind,
                    worship_token=worship_row.token,
                    source_state=state,
                    original_literal=original,
                    normalized_tokens=tokens,
                )
            )
        return ParsedProjectionAssignmentWorkbook(worship=worship, rows=tuple(rows))
    finally:
        source.close()


def resolve_projection_destination_team(team_key=SVCA_PROJECTION_TEAM_KEY):
    try:
        team = MinistryTeam.objects.get(team_key=team_key)
    except MinistryTeam.DoesNotExist as exc:
        raise ProjectionDestinationTeamError(ProjectionDestinationTeamErrorCode.MISSING) from exc
    except MinistryTeam.MultipleObjectsReturned as exc:
        raise ProjectionDestinationTeamError(ProjectionDestinationTeamErrorCode.DUPLICATE) from exc
    if not team.is_active:
        raise ProjectionDestinationTeamError(ProjectionDestinationTeamErrorCode.INACTIVE)
    if not team.is_assignable:
        raise ProjectionDestinationTeamError(ProjectionDestinationTeamErrorCode.NON_ASSIGNABLE)
    return team


def _privacy_safe_visible_identity(membership):
    if membership.display_name:
        value = membership.display_name
    elif membership.user_id:
        value = membership.user.get_full_name() or membership.user.get_username()
    else:
        raise ProjectionMappingStateError("An active Projection membership has no visible identity.")
    value = normalize_projection_identity(value)
    if not value:
        raise ProjectionMappingStateError("An active Projection membership has no visible identity.")
    return value


def _membership_candidate(membership):
    visible = _privacy_safe_visible_identity(membership)
    return ProjectionMembershipCandidate(
        membership_id=membership.pk,
        visible_identity=visible,
        identity_digest=_identity_digest(visible),
        linked_user_id=membership.user_id,
        linked_state="linked" if membership.user_id else "display_name_only",
        linked_user_active=membership.user.is_active if membership.user_id else None,
        updated_at=membership.updated_at.isoformat(),
    )


def _current_candidates(team):
    rows = TeamMembership.objects.filter(team=team, is_active=True).select_related("user").order_by("id")
    return tuple(_membership_candidate(item) for item in rows)


def _token_reviews(parsed, candidates):
    counts = Counter(token for row in parsed.rows for token in row.normalized_tokens if row.source_state == ProjectionSourceState.SUPPORTED_LITERAL)
    first_seen = []
    for row in parsed.rows:
        if row.source_state != ProjectionSourceState.SUPPORTED_LITERAL:
            continue
        for token in row.normalized_tokens:
            if token not in first_seen:
                first_seen.append(token)
    reviews = []
    for index, token in enumerate(first_seen):
        exact_ids = tuple(c.membership_id for c in candidates if normalize_projection_identity(c.visible_identity) == token)
        if len(exact_ids) == 1:
            state, prefill = ProjectionIdentityState.EXACT_PREFILL_AVAILABLE, exact_ids[0]
        elif len(exact_ids) > 1:
            state, prefill = ProjectionIdentityState.AMBIGUOUS, None
        else:
            state, prefill = ProjectionIdentityState.MAPPING_REQUIRED, None
        reviews.append(ProjectionTokenReview(index, token, counts[token], exact_ids, state, prefill))
    return tuple(reviews)


def _source_row_payload(row):
    return {
        "source_row": row.source_row,
        "source_cell": row.source_cell,
        "local_date": row.local_date.isoformat(),
        "date_kind": row.date_kind,
        "worship_token": row.worship_token,
        "source_state": row.source_state.value,
        "original_literal": row.original_literal,
        "normalized_tokens": list(row.normalized_tokens),
    }


def _candidate_payload(candidate):
    return {
        "membership_id": candidate.membership_id,
        "visible_identity": candidate.visible_identity,
        "identity_digest": candidate.identity_digest,
        "linked_user_id": candidate.linked_user_id,
        "linked_state": candidate.linked_state,
        "linked_user_active": candidate.linked_user_active,
        "updated_at": candidate.updated_at,
    }


def _team_payload(team):
    return {
        "id": team.pk,
        "key": team.team_key,
        "active": team.is_active,
        "assignable": team.is_assignable,
        "updated_at": team.updated_at.isoformat(),
    }


def prepare_projection_assignment_mapping(*, content, filename, user):
    parsed = parse_known_projection_assignment_workbook(content, filename=filename)
    profile = resolve_target_service_profile()
    team = resolve_projection_destination_team()
    candidates = _current_candidates(team)
    payload = {
        "contract_revision": MAPPING_STATE_CONTRACT_REVISION,
        "signing_version": MAPPING_SIGNING_VERSION,
        "state_type": "projection_assignment_mapping",
        "source_contract_revision": SOURCE_CONTRACT_REVISION,
        "worship_parser_contract_revision": WORSHIP_PARSER_CONTRACT_REVISION,
        "integration_key": INTEGRATION_KEY,
        "source_semantic": SOURCE_SEMANTIC,
        "source_column": SOURCE_COLUMN,
        "generated_at": timezone.now().isoformat(),
        "user_id": user.pk,
        "filename": parsed.filename,
        "sha256": parsed.sha256,
        "profile": {"id": profile.pk, "key": profile.key, "event_type": profile.event_type},
        "team": _team_payload(team),
        "rows": [_source_row_payload(row) for row in parsed.rows],
        "candidates": [_candidate_payload(item) for item in candidates],
    }
    signed = signing.dumps(payload, compress=True, salt=MAPPING_SIGNING_SALT)
    return ProjectionMappingReview(parsed, team, candidates, _token_reviews(parsed, candidates), signed)


def _safe_filename(value):
    return str(value or "workbook.xlsx").replace("\\", "/").rsplit("/", 1)[-1]


def decode_projection_assignment_mapping(token, *, user, max_age=SIGNING_MAX_AGE_SECONDS):
    try:
        payload = signing.loads(token, salt=MAPPING_SIGNING_SALT, max_age=max_age)
    except signing.BadSignature as exc:
        raise ProjectionMappingStateError("Projection mapping is invalid or expired.") from exc
    required = {
        "contract_revision", "signing_version", "state_type", "source_contract_revision",
        "worship_parser_contract_revision", "integration_key", "source_semantic", "source_column",
        "generated_at", "user_id", "filename", "sha256", "profile", "team", "rows", "candidates",
    }
    if (
        not isinstance(payload, dict)
        or set(payload) != required
        or payload["contract_revision"] != MAPPING_STATE_CONTRACT_REVISION
        or payload["signing_version"] != MAPPING_SIGNING_VERSION
        or payload["state_type"] != "projection_assignment_mapping"
        or payload["source_contract_revision"] != SOURCE_CONTRACT_REVISION
        or payload["worship_parser_contract_revision"] != WORSHIP_PARSER_CONTRACT_REVISION
        or payload["integration_key"] != INTEGRATION_KEY
        or payload["source_semantic"] != SOURCE_SEMANTIC
        or payload["source_column"] != SOURCE_COLUMN
        or payload["user_id"] != getattr(user, "pk", None)
        or not isinstance(payload["generated_at"], str)
        or not isinstance(payload["filename"], str)
        or _safe_filename(payload["filename"]) != payload["filename"]
        or not isinstance(payload["sha256"], str)
        or _SHA256_RE.fullmatch(payload["sha256"]) is None
        or not isinstance(payload["rows"], list)
        or len(payload["rows"]) != len(SUPPORTED_ROWS)
        or not isinstance(payload["candidates"], list)
    ):
        raise ProjectionMappingStateError("Projection mapping is malformed.")
    profile = resolve_target_service_profile()
    if payload["profile"] != {"id": profile.pk, "key": profile.key, "event_type": profile.event_type}:
        raise ProjectionMappingStateError("The target Service Profile changed.")
    team = resolve_projection_destination_team()
    if payload["team"] != _team_payload(team):
        raise ProjectionMappingStateError("The configured Projection team changed.")
    candidates = _current_candidates(team)
    if payload["candidates"] != [_candidate_payload(item) for item in candidates]:
        raise ProjectionMappingStateError("Projection membership identity changed.")

    rows = []
    row_keys = {"source_row", "source_cell", "local_date", "date_kind", "worship_token", "source_state", "original_literal", "normalized_tokens"}
    try:
        for item in payload["rows"]:
            if not isinstance(item, dict) or set(item) != row_keys:
                raise ValueError
            rows.append(
                ProjectionSourceRow(
                    source_row=item["source_row"],
                    source_cell=item["source_cell"],
                    local_date=date.fromisoformat(item["local_date"]),
                    date_kind=item["date_kind"],
                    worship_token=item["worship_token"],
                    source_state=ProjectionSourceState(item["source_state"]),
                    original_literal=item["original_literal"],
                    normalized_tokens=tuple(item["normalized_tokens"]),
                )
            )
    except (KeyError, TypeError, ValueError) as exc:
        raise ProjectionMappingStateError("Projection mapping rows are malformed.") from exc
    if tuple(row.source_row for row in rows) != SUPPORTED_ROWS:
        raise ProjectionMappingStateError("Projection mapping rows are not canonical.")
    for row in rows:
        index = 0 if row.source_row == 4 else row.source_row - 5
        expected_date = date(2026, 1, 4) + timedelta(weeks=index)
        if (
            type(row.source_row) is not int
            or row.source_cell != f"{SOURCE_COLUMN}{row.source_row}"
            or row.local_date != expected_date
            or row.date_kind != ("literal" if row.source_row == 4 else "formula_cached")
            or row.worship_token not in TOKEN_ORDER
            or row.source_state == ProjectionSourceState.SUPPORTED_LITERAL
            and (
                len(row.normalized_tokens) not in (1, 2)
                or len(set(row.normalized_tokens)) != len(row.normalized_tokens)
                or any(
                    normalize_projection_identity(value) != value
                    or _token_has_unsupported_syntax(value)
                    for value in row.normalized_tokens
                )
            )
            or row.original_literal is not None and not isinstance(row.original_literal, str)
        ):
            raise ProjectionMappingStateError("Projection mapping rows are not canonical.")
        if row.source_state in {
            ProjectionSourceState.SUPPORTED_LITERAL,
            ProjectionSourceState.UNSUPPORTED_TOKEN,
        } and row.original_literal is not None:
            expected = _classify_projection_cell(
                SimpleNamespace(data_type="s", value=row.original_literal)
            )
            if expected != (
                row.source_state,
                row.original_literal,
                row.normalized_tokens,
            ):
                raise ProjectionMappingStateError(
                    "Projection mapping rows are not canonical."
                )
        elif row.source_state == ProjectionSourceState.NO_SOURCE_PROPOSAL:
            if row.normalized_tokens or (
                row.original_literal is not None
                and normalize_projection_identity(row.original_literal)
            ):
                raise ProjectionMappingStateError(
                    "Projection mapping rows are not canonical."
                )
        elif row.source_state in {
            ProjectionSourceState.FORMULA_BLOCKED,
            ProjectionSourceState.ERROR_BLOCKED,
        }:
            if row.original_literal is not None or row.normalized_tokens:
                raise ProjectionMappingStateError(
                    "Projection mapping rows are not canonical."
                )
        elif row.normalized_tokens:
            raise ProjectionMappingStateError(
                "Projection mapping rows are not canonical."
            )

    from .worship_xlsx_preview import ParsedWorkbookRow

    worship_rows = tuple(
        ParsedWorkbookRow(row.source_row, f"A{row.source_row}/B{row.source_row}", row.local_date, row.date_kind, row.worship_token)
        for row in rows
    )
    worship = ParsedWorshipWorkbook(
        filename=payload["filename"],
        sha256=payload["sha256"],
        rows=worship_rows,
        token_counts=dict(Counter(row.worship_token for row in rows)),
        unsupported_rows=(
            {"row": 5, "kind": "friday", "date": "2026-01-09"},
            {"row": 57, "kind": "spillover", "date": "2027-01-03"},
            {"row": 58, "kind": "spillover", "date": "2027-01-10"},
        ),
    )
    parsed = ParsedProjectionAssignmentWorkbook(worship, tuple(rows))
    return ProjectionMappingReview(parsed, team, candidates, _token_reviews(parsed, candidates), token)


def _validate_selected_mapping(review, selected_mapping):
    candidate_by_id = {item.membership_id: item for item in review.candidates}
    allowed_tokens = {item.token for item in review.token_reviews}
    if set(selected_mapping) - allowed_tokens:
        raise ProjectionMappingValidationError("Projection mapping contains an unknown token.")
    resolved = {}
    for token in allowed_tokens:
        raw_id = selected_mapping.get(token)
        if raw_id in (None, ""):
            resolved[token] = None
            continue
        try:
            membership_id = int(raw_id)
        except (TypeError, ValueError) as exc:
            raise ProjectionMappingValidationError("Projection membership ID is malformed.") from exc
        candidate = candidate_by_id.get(membership_id)
        if candidate is None:
            raise ProjectionMappingValidationError("Selected membership is missing, inactive, or on another team.")
        resolved[token] = candidate
    return resolved


def _assignment_queryset(event_ids, team):
    members = TeamAssignmentMember.objects.select_related("membership", "membership__user").order_by("id")
    return (
        TeamAssignment.objects.filter(service_event_id__in=event_ids, ministry_team=team)
        .select_related("service_event", "ministry_team")
        .prefetch_related(Prefetch("assignment_members", queryset=members))
        .order_by("id")
    )


def _assignment_baseline(assignments):
    return [
        {
            "id": assignment.pk,
            "status": assignment.status,
            "updated_at": assignment.updated_at.isoformat(),
            "reviewed_worship_context_fingerprint": assignment.reviewed_worship_context_fingerprint,
            "members": [
                {
                    "assignment_member_id": member.pk,
                    "membership_id": member.membership_id,
                    "membership_active": member.membership.is_active,
                    "membership_updated_at": member.membership.updated_at.isoformat(),
                    "confirmed": member.confirmed_at is not None,
                    "confirmation_note_present": bool(member.confirmation_note),
                }
                for member in assignment.assignment_members.all()
            ],
        }
        for assignment in assignments
    ]


def _classify_assignment(assignments, selected_membership_ids):
    current = [item for item in assignments if item.status in CURRENT_ASSIGNMENT_STATUSES]
    historical = [item for item in assignments if item.status in HISTORICAL_ASSIGNMENT_STATUSES]
    unknown = [item for item in assignments if item.status not in {*CURRENT_ASSIGNMENT_STATUSES, *HISTORICAL_ASSIGNMENT_STATUSES}]
    current_ids = tuple(item.pk for item in current)
    historical_ids = tuple(item.pk for item in historical)
    roster = ()
    if len(current) == 1:
        roster = tuple(sorted(member.membership_id for member in current[0].assignment_members.all()))
    if len(current) > 1:
        state = ProjectionTargetState.DUPLICATE_ASSIGNMENT_BLOCKER
    elif unknown:
        state = ProjectionTargetState.UNKNOWN_ASSIGNMENT_BLOCKER
    elif historical:
        state = ProjectionTargetState.HISTORICAL_ASSIGNMENT_BLOCKER
    elif len(current) == 1:
        assignment = current[0]
        members = list(assignment.assignment_members.all())
        fingerprint = assignment.reviewed_worship_context_fingerprint
        if fingerprint is not None and WORSHIP_CONTEXT_FINGERPRINT_RE.fullmatch(fingerprint) is None:
            state = ProjectionTargetState.INVALID_ASSIGNMENT_BLOCKER
        elif any(not member.membership.is_active for member in members):
            state = ProjectionTargetState.EXISTING_ROSTER_BLOCKER
        elif len(members) == len(selected_membership_ids) and set(roster) == set(selected_membership_ids):
            state = ProjectionTargetState.EXACT_NOOP
        else:
            state = ProjectionTargetState.EXISTING_ROSTER_BLOCKER
    else:
        state = ProjectionTargetState.CREATE_CANDIDATE
    return state, current_ids, roster, historical_ids


def _target_detail(state):
    return {
        TargetMatchState.NO_TARGET: "The exact target event is missing.",
        TargetMatchState.MULTIPLE_EXACT_TARGETS: "Multiple exact target events exist.",
        TargetMatchState.TARGET_EVENT_PROFILE_FK_MISSING: "The target event has no canonical Service Profile FK.",
        TargetMatchState.TARGET_EVENT_PROFILE_IDENTITY_DRIFT: "The target event Service Profile identity drifted.",
        TargetMatchState.TARGET_EVENT_OWNED_BY_OTHER_PROFILE: "The target event belongs to another Service Profile.",
        TargetMatchState.LIFECYCLE_CONFLICT: "The target event lifecycle is unsupported.",
        TargetMatchState.AUDIENCE_INVALID_CONFLICT: "The target event audience is not ready.",
    }.get(state)


def _event_is_historical(event, *, now):
    return event.status == ServiceEvent.STATUS_COMPLETED or service_event_is_history(event, now=now)


def _membership_is_outside_event_audience(candidate, event):
    if not candidate.linked_user_id or not candidate.linked_user_active:
        return False
    audience_units = list(event.get_audience_scope_units())
    return bool(audience_units) and not user_matches_structure_audience(
        get_user_model()._default_manager.get(pk=candidate.linked_user_id), audience_units
    )


def _audience_payload(event):
    evidence = service_event_audience_readiness(event)
    return {
        "ready": evidence["ready"],
        "invalid_reasons": evidence["invalid_reasons"],
        "units": [
            {"id": unit["id"], "is_active": unit["is_active"]}
            for unit in evidence["units"]
        ],
    }


def _event_payload(event):
    local_start = timezone.localtime(event.start_datetime)
    return {
        "id": event.pk,
        "service_profile_id": event.service_profile_id,
        "expected_profile_key": event.service_profile.key,
        "scheduling_revision": event.scheduling_revision,
        "local_date": local_start.date().isoformat(),
        "local_time": local_start.time().replace(tzinfo=None).isoformat(),
        "event_type": event.event_type,
        "status": event.status,
        "audience": _audience_payload(event),
    }


def _selected_membership_payload(candidate, team):
    return {
        "id": candidate.membership_id,
        "team_id": team.pk,
        "active": True,
        "updated_at": candidate.updated_at,
        "linked_user_id": candidate.linked_user_id,
        "linked_user_active": candidate.linked_user_active,
        "visible_identity_digest": candidate.identity_digest,
    }


def build_projection_assignment_preview(*, mapping_review, selected_mapping, user, now=None):
    if not user_can_preview_projection_assignments(user):
        raise ProjectionMappingValidationError("Projection assignment preview is staff-only.")
    team = resolve_projection_destination_team()
    if team.pk != mapping_review.team.pk or _team_payload(team) != _team_payload(mapping_review.team):
        raise ProjectionMappingValidationError("The configured Projection team changed.")
    if mapping_review.candidates != _current_candidates(team):
        raise ProjectionMappingValidationError("Projection membership identity changed.")
    resolved = _validate_selected_mapping(mapping_review, selected_mapping)
    classification_now = now or timezone.now()
    matches = match_exact_service_event_targets(mapping_review.parsed.worship)
    event_ids = [match.event.pk for match in matches if match.event is not None]
    assignments_by_event = defaultdict(list)
    for assignment in _assignment_queryset(event_ids, team):
        assignments_by_event[assignment.service_event_id].append(assignment)
    token_review_by_token = {item.token: item for item in mapping_review.token_reviews}

    rows = []
    signed_rows = []
    for source, match in zip(mapping_review.parsed.rows, matches, strict=True):
        selected = tuple(resolved.get(token) for token in source.normalized_tokens)
        selected_complete = bool(selected) and all(selected)
        selected_candidates = tuple(item for item in selected if item is not None)
        event = match.event
        current_ids = ()
        roster = ()
        historical_ids = ()
        current_memberships = ()
        detail = None

        if source.source_state == ProjectionSourceState.NO_SOURCE_PROPOSAL:
            identity_state = ProjectionIdentityState.NOT_APPLICABLE
            target_state = ProjectionTargetState.NO_SOURCE_PROPOSAL
        elif source.source_state != ProjectionSourceState.SUPPORTED_LITERAL:
            identity_state = ProjectionIdentityState.NOT_APPLICABLE
            target_state = ProjectionTargetState.SOURCE_BLOCKER
            detail = "Column E is outside the exact one-or-two-person slash contract."
        elif not selected_complete:
            states = [token_review_by_token[token].identity_state for token in source.normalized_tokens if resolved.get(token) is None]
            identity_state = states[0] if len(states) == 1 else ProjectionIdentityState.MAPPING_REQUIRED
            target_state = ProjectionTargetState.IDENTITY_BLOCKER
            detail = "Every source person requires an explicit active Projection-team membership review."
        elif len({item.membership_id for item in selected_candidates}) != len(selected_candidates):
            identity_state = ProjectionIdentityState.REVIEWED_SELECTED
            target_state = ProjectionTargetState.IDENTITY_BLOCKER
            detail = "Two source people cannot resolve to the same Projection membership."
        elif match.state != TargetMatchState.EXACT_TARGET_MATCHED:
            identity_state = ProjectionIdentityState.REVIEWED_SELECTED
            target_state = ProjectionTargetState.INVALID_TARGET_BLOCKER
            detail = _target_detail(match.state)
        else:
            identity_state = ProjectionIdentityState.REVIEWED_SELECTED
            assignments = assignments_by_event[event.pk]
            selected_ids = tuple(item.membership_id for item in selected_candidates)
            assignment_state, current_ids, roster, historical_ids = _classify_assignment(assignments, selected_ids)
            current = [item for item in assignments if item.status in CURRENT_ASSIGNMENT_STATUSES]
            if len(current) == 1:
                current_memberships = tuple(_membership_candidate(member.membership) for member in current[0].assignment_members.all())
            if _event_is_historical(event, now=classification_now):
                target_state = ProjectionTargetState.HISTORICAL_EVENT_BLOCKER
                detail = "The ServiceEvent is historical; this preview does not backfill assignments."
            elif any(_membership_is_outside_event_audience(item, event) for item in selected_candidates):
                target_state = ProjectionTargetState.AUDIENCE_SAFETY_BLOCKER
                detail = "A linked selected user is outside the event audience."
            else:
                target_state = assignment_state
            signed_rows.append(
                {
                    "source": {
                        "row": source.source_row,
                        "cell": source.source_cell,
                        "semantic": SOURCE_SEMANTIC,
                        "literal_classification": source.source_state.value,
                        "token_digests": [_identity_digest(token) for token in source.normalized_tokens],
                    },
                    "event": _event_payload(event),
                    "team": _team_payload(team),
                    "memberships": [
                        {
                            **_selected_membership_payload(item, team),
                            "source_token_digest": _identity_digest(token),
                        }
                        for token, item in zip(
                            source.normalized_tokens,
                            selected_candidates,
                            strict=True,
                        )
                    ],
                    "target_state": target_state.value,
                    "assignment_baseline": _assignment_baseline(assignments),
                }
            )
        rows.append(
            ProjectionAssignmentPreviewRow(
                source, identity_state, selected_candidates, event, target_state, detail,
                current_ids, roster, historical_ids, current_memberships,
            )
        )

    profile = resolve_target_service_profile()
    payload = {
        "contract_revision": PREVIEW_CONTRACT_REVISION,
        "signing_version": PREVIEW_SIGNING_VERSION,
        "state_type": "projection_assignment_preview",
        "source_contract_revision": SOURCE_CONTRACT_REVISION,
        "worship_parser_contract_revision": WORSHIP_PARSER_CONTRACT_REVISION,
        "integration_key": INTEGRATION_KEY,
        "source_semantic": SOURCE_SEMANTIC,
        "source_column": SOURCE_COLUMN,
        "generated_at": timezone.now().isoformat(),
        "user_id": user.pk,
        "filename": mapping_review.parsed.filename,
        "sha256": mapping_review.parsed.sha256,
        "profile": {"id": profile.pk, "key": profile.key, "event_type": profile.event_type},
        "team": _team_payload(team),
        "rows": signed_rows,
    }
    signed = signing.dumps(payload, compress=True, salt=PREVIEW_SIGNING_SALT)
    return ProjectionAssignmentPreview(mapping_review, tuple(rows), payload, signed)


def decode_signed_projection_assignment_preview(token, *, user, max_age=SIGNING_MAX_AGE_SECONDS, now=None):
    """Decode and revalidate signed zero-write evidence; never grants write authority."""

    try:
        payload = signing.loads(token, salt=PREVIEW_SIGNING_SALT, max_age=max_age)
    except signing.BadSignature as exc:
        raise SignedProjectionPreviewError("Signed Projection preview is invalid or expired.") from exc
    required = {
        "contract_revision", "signing_version", "state_type", "source_contract_revision",
        "worship_parser_contract_revision", "integration_key", "source_semantic", "source_column",
        "generated_at", "user_id", "filename", "sha256", "profile", "team", "rows",
    }
    if (
        not isinstance(payload, dict)
        or set(payload) != required
        or payload["contract_revision"] != PREVIEW_CONTRACT_REVISION
        or payload["signing_version"] != PREVIEW_SIGNING_VERSION
        or payload["state_type"] != "projection_assignment_preview"
        or payload["source_contract_revision"] != SOURCE_CONTRACT_REVISION
        or payload["worship_parser_contract_revision"] != WORSHIP_PARSER_CONTRACT_REVISION
        or payload["integration_key"] != INTEGRATION_KEY
        or payload["source_semantic"] != SOURCE_SEMANTIC
        or payload["source_column"] != SOURCE_COLUMN
        or payload["user_id"] != getattr(user, "pk", None)
        or not isinstance(payload["generated_at"], str)
        or not isinstance(payload["filename"], str)
        or _safe_filename(payload["filename"]) != payload["filename"]
        or not isinstance(payload["sha256"], str)
        or _SHA256_RE.fullmatch(payload["sha256"]) is None
        or not isinstance(payload["rows"], list)
    ):
        raise SignedProjectionPreviewError("Signed Projection preview is malformed.")
    profile = resolve_target_service_profile()
    if payload["profile"] != {"id": profile.pk, "key": profile.key, "event_type": profile.event_type}:
        raise SignedProjectionPreviewError("The target Service Profile changed.")
    try:
        team = resolve_projection_destination_team()
    except ProjectionDestinationTeamError as exc:
        raise SignedProjectionPreviewError(
            "The configured Projection team changed."
        ) from exc
    if payload["team"] != _team_payload(team):
        raise SignedProjectionPreviewError("The configured Projection team changed.")

    classification_now = now or timezone.now()
    seen_cells = set()
    for row in payload["rows"]:
        try:
            source = row["source"]
            event_state = row["event"]
            membership_states = row["memberships"]
            if (
                set(row) != {"source", "event", "team", "memberships", "target_state", "assignment_baseline"}
                or set(source) != {"row", "cell", "semantic", "literal_classification", "token_digests"}
                or source["semantic"] != SOURCE_SEMANTIC
                or source["literal_classification"] != ProjectionSourceState.SUPPORTED_LITERAL.value
                or source["cell"] != f"{SOURCE_COLUMN}{source['row']}"
                or source["row"] not in SUPPORTED_ROWS
                or source["cell"] in seen_cells
                or not isinstance(source["token_digests"], list)
                or len(source["token_digests"]) not in (1, 2)
                or any(not isinstance(value, str) or _SHA256_RE.fullmatch(value) is None for value in source["token_digests"])
                or row["team"] != _team_payload(team)
                or not isinstance(membership_states, list)
                or len(membership_states) != len(source["token_digests"])
            ):
                raise ValueError
            seen_cells.add(source["cell"])
            source_index = 0 if source["row"] == 4 else source["row"] - 5
            expected_local_date = date(2026, 1, 4) + timedelta(
                weeks=source_index
            )
            event = ServiceEvent.objects.select_related("service_profile").prefetch_related("audience_scope_links__unit").get(pk=event_state["id"])
            local_start = timezone.localtime(event.start_datetime)
            identity = inspect_service_profile_identity(event)
            if (
                event_state != _event_payload(event)
                or not identity.is_exact
                or event.service_profile_id != profile.pk
                or event.service_profile.key != SUPPORTED_PROFILE_KEY
                or event.event_type != SUPPORTED_EVENT_TYPE
                or event_state["local_date"] != expected_local_date.isoformat()
                or local_start.time().replace(tzinfo=None) != SUPPORTED_LOCAL_TIME
                or event.status not in {ServiceEvent.STATUS_PUBLISHED, ServiceEvent.STATUS_COMPLETED}
                or not service_event_audience_readiness(event)["ready"]
            ):
                raise ValueError
            candidates = []
            for token_digest, membership_state in zip(
                source["token_digests"], membership_states, strict=True
            ):
                membership = TeamMembership.objects.select_related("user").get(pk=membership_state["id"])
                candidate = _membership_candidate(membership)
                expected_membership_state = {
                    **_selected_membership_payload(candidate, team),
                    "source_token_digest": token_digest,
                }
                if (
                    not membership.is_active
                    or membership.team_id != team.pk
                    or membership_state != expected_membership_state
                ):
                    raise ValueError
                candidates.append(candidate)
            selected_ids = tuple(item.membership_id for item in candidates)
            if len(set(selected_ids)) != len(selected_ids):
                raise ValueError
            assignments = list(_assignment_queryset((event.pk,), team))
            target_state, _current_ids, _roster, _history = _classify_assignment(assignments, selected_ids)
            if _event_is_historical(event, now=classification_now):
                target_state = ProjectionTargetState.HISTORICAL_EVENT_BLOCKER
            elif any(_membership_is_outside_event_audience(item, event) for item in candidates):
                target_state = ProjectionTargetState.AUDIENCE_SAFETY_BLOCKER
            if row["target_state"] != target_state.value or row["assignment_baseline"] != _assignment_baseline(assignments):
                raise ValueError
        except (
            KeyError, TypeError, ValueError, ServiceEvent.DoesNotExist,
            TeamMembership.DoesNotExist, ProjectionMappingStateError,
        ) as exc:
            raise SignedProjectionPreviewError("Signed Projection preview is stale or malformed.") from exc
    return payload
